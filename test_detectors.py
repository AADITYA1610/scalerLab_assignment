"""Constructed checks for rare PII and previously observed boundary errors."""
from docx import Document
from redact import detect, find_matches, gather_values

cases = [
    ("SSN: 123-45-6789", "SSN", "123-45-6789"),
    ("Card: 4111 1111 1111 1111", "CREDIT_CARD", "4111 1111 1111 1111"),
    ("Date of Birth: 12 March 1985", "DATE_OF_BIRTH", "12 March 1985"),
    ("IP: 192.168.1.12", "IP_ADDRESS", "192.168.1.12"),
    ("Offer closes 12 March 1985", None, None),
    ("Card: 4111 1111 1111 1112", None, None),
]

for text, expected_kind, expected_value in cases:
    actual = [(m.kind, text[m.start:m.end]) for m in detect(text)]
    expected = [] if expected_kind is None else [(expected_kind, expected_value)]
    assert actual == expected, f"{text!r}: expected {expected}, got {actual}"

boundary_cases = [
    (
        "Nuvama Wealth Management Limited and ICICI Securities Limited",
        ["Nuvama Wealth Management Limited", "ICICI Securities Limited"],
    ),
    ("KSH Infra Park 5 Private Limited", ["KSH Infra Park 5 Private Limited"]),
    ("Self-Certified Syndicate Bank", []),
    (
        "SEBI Bhavan, Plot No. C4 A, ‘G’ Block",
        ["SEBI Bhavan, Plot No. C4 A, ‘G’ Block"],
    ),
]
for value, expected in boundary_cases:
    actual = [value[m.start:m.end] for m in find_matches(value, [])]
    assert actual == expected, f"{value!r}: expected {expected}, got {actual}"

doc = Document()
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Name"
table.cell(0, 1).text = "Designation"
table.cell(1, 0).text = "Katyayani Balasubramanian"
table.cell(1, 1).text = "Independent Director"
assert "Katyayani Balasubramanian" in gather_values(doc)["PERSON"]

print(f"Passed {len(cases) + len(boundary_cases) + 1} constructed detector checks")
